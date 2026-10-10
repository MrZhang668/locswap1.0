"""Write the review workbook for the 530-query set.

    python scripts/intent_eval/build_excel.py

Five sheets, in the order a reviewer reads them:
  00_数据口径        how the set was built and what it cannot tell you
  01_Query总表       the 530 rows, one message per row
  02_意图分类表      the taxonomy with volumes and what each leaf is testing
  03_覆盖度统计      the coverage pivots (scenario / style / category / quality)
  04_待人工裁定      the rows whose gold label the prompts do not settle
"""

from __future__ import annotations

import collections
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

import dataset  # noqa: E402
from taxonomy import (  # noqa: E402
    CATEGORIES, FAMILIES, LEAVES, SCORED_SEPARATELY, SELLER_PROFILES, STYLES,
)

OUT = Path(__file__).resolve().parents[2] / "docs" / "intent_eval" / "LocSwap_卖家意图_500query_v2_中文.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=12, color="1F3864")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header(ws, headers: list[str], widths: list[int], row: int = 1) -> None:
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _write(ws, r: int, values: list, wrap_cols: tuple[int, ...] = ()) -> None:
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = BORDER
        cell.alignment = Alignment(
            vertical="top", wrap_text=c in wrap_cols,
            horizontal="left" if c in wrap_cols else "center",
        )


# ------------------------------------------------------------------
def sheet_method(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("00_数据口径")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 118
    lines = [
        ("评测对象", "LocSwap 卖家顾问入口 POST /api/seller/advisor/chat。买家搜索入口 /api/search/chat 不在本次范围内。"),
        ("要回答的两个问题",
         "① 意图识别准确度：路由器把一条卖家消息判给哪个专家（sales_agent / publishing_agent / clarify）、"
         "判成哪个分析场景（price / demand / both / find_buyers / none）、抽出哪些槽位。"
         "② 模型是否能触发：判对之后，下游是否真的动起来了——该调的工具有没有调、"
         "发布合规流水线有没有跑、推送写操作的确认闸门有没有按预期开合。"),
        ("样本量", "530 条 = 原 500（A-E 族）+ 新增 30（F 跨agent多意图）。"
                "原 500 的构造分布、qid 与覆盖度统计一律未动，两次运行之间可比。"),
        ("计数口径",
         "一条消息 = 一行。不按用户 ID 去重，不按会话分组，同一个卖家问两次相同问题算两行。"
         "40 行带前置轮次（指代消解与推送确认），这些行会真实发两轮，但仍只按被测的那一条计一行。"),
        ("数据来源与其局限",
         "平台目前没有 query 日志表（docs/buyer_intent_session_log_plan.md 里的 SearchSession 仍是待实施状态），"
         "因此不存在真实的「近一周」可提取数据。这 500 条是按覆盖度规格构造的，用于替代一周流量。"
         "代价是分布不是观测出来的；收益是每个意图叶子、每种表达风格、每种卖家状态都有确定的样本数，"
         "而不是碰运气。报告中所有占比都应读作「构造分布」而非「实测流量分布」。"),
        ("F 族为什么单列",
         "F 族考的是一条消息同时需要两个专家：把东西挂上去，并且告诉我它值多少 / 谁想要。"
         "它的「金标-路由序列」写的是产品意义上应该发生什么（publishing_agent -> sales_agent），"
         "而不是成文提示词推导出的结果——当前架构一轮一路由，图里 sales 与 publishing 之间没有边，"
         "而且三个分析工具的 item_id 都要先 resolve_owned_item，商品必须已落库（中间隔着人工发布表单）。"
         "所以这条序列今天不可达。把它并进主准确率会把架构缺口记成模型判错，因此 F 族不计入 L1/L2 分母，"
         "单独报两个数：第一跳路由准确率，以及多意图保全率（第二个意图被应答 / 被承认 / 被静默丢弃）。"),
        ("金标-必触发工具",
         "每个叶子声明两组工具：必触发（少一个就算漏触发）与容许（出现不算误触发，"
         "典型是 list_my_items——卖家用「那辆车」指代时，调它来解析 item_ref 是合法的）。"
         "两组之外的工具出现即记一次误触发。A7 的必触发是空集：卖家的在售清单已由 _listings_context "
         "内联进提示词，提示词明说已有的就别再调，所以调了不算错、不调也不算漏。"),
        ("量级分配依据",
         "原 500：A 销售分析 41% / B 挂售发布 28% / C 需澄清 14% / D 边界与风险 12% / E 动作闸门 5%。"
         "依据是卖家顾问的产品定位：它主要服务已经有在售商品的卖家，所以销售分析占大头；"
         "挂售是第二大入口；澄清与风险类按二手平台 IM 的常见噪声比例给量。这是一个假设，"
         "上线真实日志后应当用实测分布重新校准。"),
        ("商品品类覆盖",
         "9 个平台类目全覆盖，类目名取自切中文后的 CATEGORY_DICT（电子产品 / 自行车与滑板车 / 家具与电器 / "
         "图书与教材 / 服装与鞋靴 / 运动与户外 / 家居与厨房 / 乐器 / 学习用品）。"
         "A 族（销售分析）的品类受限于测试卖家的真实在售商品——只能问他确实挂着的东西，"
         "否则问题本身就不成立；品类广度主要由 B 族承担。电子产品偏多是校园二手的真实形态，不做人为拉平。"),
        ("语言口径",
         "平台已切换为中文：prompts/ 下的 4 个运行时提示词、CATEGORY_DICT、CONDITION_DICT、"
         "listing_rules 的规则标题、以及测试卖家的商品名均为中文（规则码 BAN-*/WARN-* 仍是英文标识符）。"
         "数据集主体因此为中文 query。其中 22 条刻意写成英文，标记为 non_english——"
         "在中文产品里，英文输入才是「非产品语言」的鲁棒性用例，方向与英文版时相反。"),
        ("表达覆盖", "12 种表达风格，见 03_覆盖度统计。含详细 / 简洁 / 清晰 / 含糊 / 长尾 / 歧义 / 矛盾 / "
                 "错拼 / 指代 / 多意图 / 非产品语言 / 噪声。"),
        ("好 query / 差 query", "quality 列标注。good = 一个称职的人类看了能直接动手；bad = 看了也不知道要干什么。"),
        ("卖家状态", " | ".join(f"{k}：{v['desc']}" for k, v in SELLER_PROFILES.items())
         + "。路由结论是卖家状态的函数（has_active_items 由后端注入），所以状态必须是数据集的一部分。"),
        ("金标来源", "先按 prompts/orchestrator_router.md 的成文规则机器推导，再人工复核。"
                 "label_conf=review 的行表示成文规则本身没有唯一解，需要人工裁定——见 04 表。"
                 "例外是 F 族的路由序列，它是产品目标而非提示词推导结果，见「F 族为什么单列」。"),
        ("v2 修订：5 行不可达金标",
         "Q037 / Q071 / Q106 / Q168 的金标原为 publishing_agent 配 price / demand / find_buyers，"
         "Q142 原为 clarify 配 demand。这些组合系统永远产不出来——router._sanitize() 对任何非 "
         "sales_agent 的路由强制把 scenario 归零并记入 dropped_slots。也就是说这 5 行在场景维度上"
         "无论模型多好都必错。已一律改为 scenario=none，被丢掉的那个意图移到新列「金标-第二意图」，"
         "由多意图保全率去度量，不再污染场景准确率。"),
        ("长度上限", "API 的 message 字段上限 500 字符，数据集内所有 query 均未越界。"),
    ]
    ws.cell(row=1, column=1, value="LocSwap 卖家意图识别评测 · 数据口径说明").font = TITLE_FONT
    r = 3
    for k, v in lines:
        a = ws.cell(row=r, column=1, value=k)
        a.font = Font(bold=True, size=10)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=r, column=2, value=v)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = max(30, 15 * (len(v) // 60 + 1))
        r += 1
    ws.cell(row=r + 1, column=1, value="生成时间").font = Font(bold=True, size=10)
    ws.cell(row=r + 1, column=2, value=f"共 {len(rows)} 行，版本 v2（+F 族 30 条；修 5 行不可达金标；新增 4 个金标列）")


def sheet_queries(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("01_Query总表")
    heads = ["qid", "意图族", "叶子类", "叶子中文名", "query 文本", "表达风格", "商品品类",
             "query 质量", "卖家状态", "带图", "前置轮次", "金标-路由", "金标-场景",
             "金标-路由序列", "金标-必触发工具", "金标-容许工具", "金标-第二意图",
             "标注置信", "备注"]
    widths = [8, 14, 7, 18, 62, 12, 18, 9, 10, 8, 34, 16, 12,
              32, 34, 22, 14, 10, 42]
    _header(ws, heads, widths)
    for i, r in enumerate(rows, start=2):
        _write(ws, i, [
            r["qid"], r["family"], r["leaf"], r["leaf_cn"], r["text"],
            f'{r["style"]} {STYLES[r["style"]].split("：")[0]}', r["cat"] or "—",
            r["quality"], f'{r["seller"]}({SELLER_PROFILES[r["seller"]]["n_items"]}件)',
            "是" if r["photo"] else "", r["ctx"] or "",
            r["gold_agent"] or "（不经路由）", r["gold_scenario"] or "—",
            r["gold_seq"] or "（不经路由）",
            " + ".join(r["tools_req"]) or "—",
            " / ".join(r["tools_ok"]) or "—",
            r["second_intent"] or "—",
            r["label_conf"], r["note"],
        ], wrap_cols=(5, 11, 19))
        if r["label_conf"] == "review":
            for c in range(1, len(heads) + 1):
                ws.cell(row=i, column=c).fill = REVIEW_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{len(rows) + 1}"


def sheet_taxonomy(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("02_意图分类表")
    heads = ["意图族", "叶子类", "中文名", "金标-路由", "金标-场景", "金标-必触发工具",
             "量级", "占比", "计入主准确率", "这一类在考什么"]
    _header(ws, heads, [16, 8, 20, 18, 13, 34, 8, 8, 14, 60])
    cnt = collections.Counter(r["leaf"] for r in rows)
    total = len(rows)
    r = 2
    for fam in FAMILIES:
        fam_n = sum(n for k, n in cnt.items() if LEAVES[k]["family"] == fam)
        for leaf, spec in LEAVES.items():
            if spec["family"] != fam:
                continue
            _write(ws, r, [fam, leaf, spec["cn"], spec["agent"] or "（不经路由）",
                           spec["scn"] or "—", " + ".join(spec["tools_req"]) or "—",
                           cnt[leaf], f"{cnt[leaf] / total:.1%}",
                           "否" if fam in SCORED_SEPARATELY else "是",
                           spec["tests"]], wrap_cols=(10,))
            r += 1
        _write(ws, r, [f"{fam} 小计", "", "", "", "", "", fam_n,
                       f"{fam_n / total:.1%}", "", ""])
        for c in range(1, 11):
            ws.cell(row=r, column=c).font = Font(bold=True)
        r += 1
    _write(ws, r, ["合计", "", "", "", "", "", total, "100.0%", "", ""])
    for c in range(1, 11):
        ws.cell(row=r, column=c).font = Font(bold=True)


def sheet_coverage(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("03_覆盖度统计")
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["H"].width = 52
    r = 1

    def block(title: str, counter: collections.Counter, note: str = "",
              describe: dict | None = None) -> None:
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = TITLE_FONT
        r += 1
        for h, c in (("取值", 1), ("条数", 2), ("占比", 3), ("说明", 4)):
            cell = ws.cell(row=r, column=c, value=h)
            cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        r += 1
        for k, n in counter.most_common():
            _write(ws, r, [str(k), n, f"{n / len(rows):.1%}",
                           (describe or {}).get(k, "")], wrap_cols=(4,))
            r += 1
        if note:
            ws.cell(row=r, column=1, value=note).font = Font(italic=True, size=9, color="808080")
            r += 1
        r += 1

    block("① 金标路由分布（L1）", collections.Counter(r_["gold_agent"] or "（不经路由 · E族）" for r_ in rows),
          "clarify 占 20.6% 是刻意的：路由器的不确定默认值就是 clarify，这一档样本不足就测不出它是否过度犹豫。")
    block("② 金标场景分布（L2，仅 sales_agent 有意义）",
          collections.Counter(r_["gold_scenario"] or "（不适用）" for r_ in rows))
    block("③ 表达风格覆盖", collections.Counter(r_["style"] for r_ in rows), describe=STYLES)
    block("④ 商品品类覆盖", collections.Counter(r_["cat"] or "（无品类：澄清/闲聊/风险类）" for r_ in rows),
          "A 族品类受测试卖家真实在售商品约束；品类广度主要由 B 族承担。")
    block("⑤ query 质量", collections.Counter(r_["quality"] for r_ in rows),
          describe={"good": "一个称职的人类看了能直接动手", "bad": "看了也不知道要干什么，必须反问"})
    block("⑥ 卖家状态", collections.Counter(r_["seller"] for r_ in rows),
          describe={k: v["desc"] for k, v in SELLER_PROFILES.items()})
    block("⑦ 标注置信度", collections.Counter(r_["label_conf"] for r_ in rows),
          describe={"high": "成文 prompt 规则可唯一推导", "review": "规则未覆盖或两解皆可，需人工裁定"})
    block("⑧ 多轮 / 带图", collections.Counter(
        ("带图（has_photo=true）" if r_["photo"] else "带前置轮次" if r_["ctx"] else "单轮纯文本")
        for r_ in rows))

    ws.cell(row=r, column=1, value="⑨ 表达风格 × 意图族 交叉").font = TITLE_FONT
    r += 1
    for c, fam in enumerate(["表达风格"] + FAMILIES + ["合计"], start=1):
        cell = ws.cell(row=r, column=c, value=fam)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    cross = collections.Counter((r_["style"], r_["family"]) for r_ in rows)
    for style in STYLES:
        vals = [cross[(style, f)] for f in FAMILIES]
        _write(ws, r, [style] + vals + [sum(vals)])
        r += 1


def sheet_review(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("04_待人工裁定")
    ws.cell(row=1, column=1,
            value="这些行的金标不是机器推导出来的——成文 prompt 对它们没有唯一解。"
                  "请逐行给出裁定；未裁定前，它们在准确率分母里但结论应单列。").font = TITLE_FONT
    heads = ["qid", "叶子类", "query 文本", "机器初判-路由", "机器初判-场景", "争议点", "人工裁定-路由", "人工裁定-场景", "裁定理由"]
    _header(ws, heads, [8, 8, 60, 18, 14, 50, 16, 14, 40], row=3)
    r = 4
    for row in rows:
        if row["label_conf"] != "review":
            continue
        _write(ws, r, [row["qid"], row["leaf"], row["text"], row["gold_agent"] or "（不经路由）",
                       row["gold_scenario"] or "—", row["note"], "", "", ""], wrap_cols=(3, 6, 9))
        r += 1
    ws.cell(row=r + 1, column=1, value=f"共 {r - 4} 行待裁定").font = Font(bold=True)


def _check_categories(rows: list[dict]) -> None:
    """Fail loudly if a row names a category the platform no longer has.

    Worth a guard rather than a comment: the platform was switched from English
    to Chinese category names mid-project, and a stale value here would travel
    into the gold labels and quietly score a correct extraction as wrong.
    """
    bad = sorted({r["cat"] for r in rows if r["cat"] and r["cat"] not in CATEGORIES})
    if bad:
        raise SystemExit(f"这些品类不在 CATEGORY_DICT 里，数据集与平台配置已经漂移: {bad}")


def main() -> None:
    rows = dataset.load()
    _check_categories(rows)
    wb = Workbook()
    wb.remove(wb.active)
    sheet_method(wb, rows)
    sheet_queries(wb, rows)
    sheet_taxonomy(wb, rows)
    sheet_coverage(wb, rows)
    sheet_review(wb, rows)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}  ({len(rows)} rows)")


if __name__ == "__main__":
    main()
